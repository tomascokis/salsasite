#!/usr/bin/env python3

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
WORKBOOK_PATH = ROOT / "data" / "legacy" / "reference" / "data_reference.xlsx"
OUTPUT_DIR = ROOT / "data" / "live" / "bootstrap" / "catalog"
RAW_MOVES_PATH = OUTPUT_DIR / "raw-moves.json"


def normalize_header(value: object, index: int) -> str:
    text = "" if value is None else str(value).strip()
    if not text:
        return f"unnamed_{index}"

    normalized = []
    for char in text:
        if char.isalnum():
            normalized.append(char.lower())
        else:
            normalized.append("_")

    key = "".join(normalized).strip("_")
    while "__" in key:
        key = key.replace("__", "_")

    return key or f"unnamed_{index}"


def cell_value(value: object):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    header_rows = list(sheet.iter_rows(min_row=1, max_row=2, values_only=True))
    if len(header_rows) < 2:
        raise RuntimeError("Expected at least two header rows in data_reference.xlsx")

    headers = [normalize_header(value, index) for index, value in enumerate(header_rows[1], start=1)]

    records = []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        if all(value is None for value in row):
            continue

        record = {headers[index]: cell_value(value) for index, value in enumerate(row)}
        if not record.get("id"):
            continue
        records.append(record)

    RAW_MOVES_PATH.write_text(json.dumps(records, indent=2, ensure_ascii=True), encoding="utf-8")

    print(f"Wrote {RAW_MOVES_PATH.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
